"""Builds web/public/media/passport-pay.pptx and .pdf for the pitch.

The deck is the official HackMeridian 2026 template filled by deck/fill-template.py (the copy
design/hackmeridian-template/PITCH.md asks to reuse). This driver runs it and then keeps the file
in sync with the rest of the project, without touching the template chrome:

  * speaker notes come from site/presentacion/index.html, the same script /notas shows (the
    honest variant until the real injection run exists, see site/evidence.js);
  * the REJECTED stamp on slide 2 carries the injection run's tx once it exists, else the real
    CapExceeded tx; the contract test count is today's real one (CONTRACT_TESTS below);
  * checks: 5 slides, no em dash, no template guidance text left, no other person named;
  * exports the PDF with LibreOffice.

Usage (python-pptx in a venv in /var/tmp, LibreOffice for the PDF):
    python3 -m venv /var/tmp/pp-venv && /var/tmp/pp-venv/bin/pip install python-pptx cssselect
    /var/tmp/pp-venv/bin/python deck/build-pptx.py            # .pptx + .pdf
    /var/tmp/pp-venv/bin/python deck/build-pptx.py --no-pdf   # only the .pptx
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import html as lhtml
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "design/hackmeridian-template/template.pptx"
FILL = ROOT / "deck/fill-template.py"
DECK_HTML = ROOT / "site/presentacion/index.html"
EVIDENCE = ROOT / "site/evidence.js"
OUT = ROOT / "web/public/media/passport-pay.pptx"
DASH = chr(0x2014)  # em dash, banned from the deck
# cargo test on 26 oct 2026: PassportAccount 51 + PassportFactory 11 (scripts/test-all.sh).
CONTRACT_TESTS = 62


def clean(s):
    return re.sub(r"\s+", " ", s or "").strip()


def load_evidence():
    src = EVIDENCE.read_text("utf8")
    return json.loads(src[src.index("{"): src.rindex("}") + 1])


def load_notes(ev):
    """Per slide: (title, start, budget, [paragraphs]) from the web deck, cues prefixed."""
    doc = lhtml.fromstring(DECK_HTML.read_text("utf8"))
    out = []
    for s in doc.cssselect("section.slide"):
        inj = s.cssselect('aside.note[data-when="injection"]')
        base = [a for a in s.cssselect("aside.note") if a.get("data-when") is None]
        note = inj[0] if (ev.get("injection") and inj) else base[0]
        paras = []
        for p in note.cssselect("p"):
            t = clean(p.text_content())
            if "cue" in (p.get("class") or "") and not t.lower().startswith("cue"):
                t = "Cue: " + t
            paras.append(t)
        out.append((s.get("data-title"), s.get("data-time"), s.get("data-dur"), paras))
    if len(out) != 5:
        raise SystemExit(f"expected 5 slides in {DECK_HTML}, found {len(out)}")
    return out


def replace_in_runs(prs, old, new):
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        n += 1
    return n


def sync(path, ev):
    prs = Presentation(path)
    # contract test count: the real one
    replace_in_runs(prs, "40+ contract tests", f"{CONTRACT_TESTS} contract tests")
    # wording kept as narrow as what is enforced today (docs/DECK.md, "Claims check")
    replace_in_runs(prs, "backed by a ZK human passport", "each with a ZK passport credential, one per agent, anti-Sybil by nullifier")
    replace_in_runs(prs, "one injected page can drain an agent's wallet.", "one injected page can drain an agent's wallet. Ours lost zero.")
    # slide 2 stamp: the injection run's tx once it exists
    inj = ev.get("injection")
    if inj:
        short = f"{inj['tx'][:8]}…{inj['tx'][-4:]}"
        cap = ev["capExceeded"]["tx"]
        replace_in_runs(prs, "CapExceeded #6", f"{inj.get('error', 'PayeeNotAllowed')} #{inj.get('code', 10)}")
        replace_in_runs(prs, f"tx {cap[:8]}…{cap[-4:]}", f"tx {short}")
        for sh in prs.slides[1].shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.hyperlink.address and cap in r.hyperlink.address:
                            r.hyperlink.address = f"{ev['explorer']}/tx/{inj['tx']}"
    # speaker notes: the /notas script, with its timing line
    for s, (title, start, dur, paras) in zip(prs.slides, load_notes(ev)):
        s.notes_slide.notes_text_frame.text = "\n".join([f"{title} · starts {start} · budget {dur} s"] + paras)
    prs.core_properties.title = "Passport Pay · HackMeridian 2026 pitch deck"
    prs.core_properties.author = "Passport Pay"
    prs.save(path)


def check(path):
    prs = Presentation(path)
    text = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                text.append(sh.text_frame.text)
        text.append(s.notes_slide.notes_text_frame.text)
    blob = "\n".join(text)
    problems = []
    if len(prs.slides) != 5:
        problems.append(f"{len(prs.slides)} slides, expected 5")
    if DASH in blob:
        problems.append("em dash in the deck")
    for guide in ("Show how your project", "Present your idea", "This is where you explain", "Here you explain", "Finally, explain", "Think:"):
        if guide in blob:
            problems.append(f"template guidance text left: {guide!r}")
    if "Leo" in blob or "Cagliero" in blob:
        problems.append("a person other than Lucas Cariddi is named")
    if "40+" in blob:
        problems.append("test count is not the real one")
    # the deployed validator does not filter who mints: the passport is not a proof of a human
    if re.search(r"\bhumans?\b", blob, re.I):
        problems.append("the deck says 'human' about the passport (use: a ZK passport credential, one per agent, anti-Sybil by nullifier)")
    if problems:
        raise SystemExit("; ".join(problems))
    return len(prs.slides)


def to_pdf(pptx, dest):
    with tempfile.TemporaryDirectory(dir="/var/tmp") as tmp:
        profile = f"-env:UserInstallation=file://{tmp}/lo"
        subprocess.run(["soffice", profile, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(pptx)], check=True, capture_output=True, timeout=240)
        shutil.copy(Path(tmp) / (pptx.stem + ".pdf"), dest)
    return dest


if __name__ == "__main__":
    ev = load_evidence()
    subprocess.run([sys.executable, str(FILL), str(TEMPLATE), str(OUT)], check=True)
    sync(OUT, ev)
    n = check(OUT)
    stamp = (ev.get("injection") or ev["capExceeded"])["tx"]
    print(f"{OUT.relative_to(ROOT)}: {n} slides, {OUT.stat().st_size} bytes, stamp tx {stamp[:8]}")
    if "--no-pdf" not in sys.argv:
        pdf = to_pdf(OUT, OUT.with_suffix(".pdf"))
        print(f"{pdf.relative_to(ROOT)}: {pdf.stat().st_size} bytes")
